from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient

from api.models import Organization, OrganizationMembership
from .kpi_import import KpiImportError, _parse_workbook
from .models import KpiDailyRecord, KpiDayState, KpiImportBatch, KpiPeriodRecord


def tracker_bytes(fiscal_year, *, target=12918, prior_metrics=None, month_plan=652449, mtd_sales=119829):
    """Build the supported 12-period, 52-week tracker shape entirely in memory."""
    from openpyxl import Workbook

    prior_metrics = prior_metrics or {}
    workbook = Workbook()
    workbook.remove(workbook.active)
    period_weeks = (4, 4, 5) * 4
    current_date = date(fiscal_year - 1, 12, 1)
    current_date += timedelta(days=(-current_date.weekday()) % 7)
    week_index = 0

    for period, number_of_weeks in enumerate(period_weeks, start=1):
        month_number = 12 if period == 1 else period - 1
        sheet_year = fiscal_year - 1 if period == 1 else fiscal_year
        month_name = date(sheet_year, month_number, 1).strftime('%b')
        sheet = workbook.create_sheet(f'{month_name} {sheet_year}')
        sheet.append([
            'Week', 'Incentive Targets', 'Day', 'Retail Sales Net Value', 'Traffic',
            'Conversion Day', 'Retail Sales Net ATV Day', 'Retail Sales Net UPT Day',
        ])

        for week_in_period in range(1, number_of_weeks + 1):
            monday = current_date + timedelta(days=week_index * 7)
            sunday = monday + timedelta(days=6)
            sheet.append([f'Wk {week_in_period}', None, None, None, None, None, None, None])
            sheet.append([f'{monday.month}/{monday.day}-{sunday.month}/{sunday.day}'])
            for weekday in range(7):
                day_date = monday + timedelta(days=weekday)
                metrics = prior_metrics.get(day_date.isoformat(), {})
                day_target = target if fiscal_year == 2026 and day_date == date(2026, 9, 14) else 10000
                sales = metrics.get('sales', 9000)
                traffic = metrics.get('traffic', 500)
                conversion = metrics.get('conversion', 0.15)
                atv = metrics.get('atv', 90)
                upt = metrics.get('upt', 1.8)
                sheet.append([
                    None, day_target, day_date.strftime('%A'), sales, traffic,
                    conversion, atv, upt,
                ])
            week_index += 1

        sheet.append([None, 'MTD', month_plan if period == 10 else 100000, mtd_sales if period == 10 else 50000])

    result = BytesIO()
    workbook.save(result)
    workbook.close()
    return result.getvalue()


def upload(name, content):
    return SimpleUploadedFile(name, content, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def mutate_tracker(content, change):
    from openpyxl import load_workbook

    source = BytesIO(content)
    workbook = load_workbook(source)
    change(workbook)
    result = BytesIO()
    workbook.save(result)
    workbook.close()
    return result.getvalue()


class KpiImportParserTests(TestCase):
    def test_normal_tracker_extracts_fiscal_period_and_daily_values(self):
        prior = tracker_bytes(
            2025,
            prior_metrics={
                '2025-09-15': {
                    'sales': 9522, 'traffic': 628, 'conversion': 0.172,
                    'atv': 88.17, 'upt': 1.76,
                },
            },
        )
        current = tracker_bytes(
            2026,
        )

        parsed = _parse_workbook(upload('tracker.xlsx', current))
        by_date = {row['business_date']: row for row in parsed.days}
        september = next(row for row in parsed.periods if row['fiscal_period'] == 10)

        self.assertEqual(parsed.fiscal_year, 2026)
        self.assertEqual(len(parsed.days), 364)
        self.assertEqual(by_date[date(2026, 9, 14)]['day_target'], 12918)
        self.assertEqual(september['sales_plan'], 652449)
        self.assertEqual(september['mtd_sales'], 119829)

        prior_parsed = _parse_workbook(upload('prior.xlsx', prior))
        comparison = {row['business_date']: row for row in prior_parsed.days}[date(2025, 9, 15)]
        self.assertEqual(comparison['sales'], 9522)
        self.assertEqual(comparison['traffic'], 628)
        self.assertEqual(comparison['conversion'], Decimal('0.172'))
        self.assertEqual(comparison['atv'], Decimal('88.17'))
        self.assertEqual(comparison['upt'], Decimal('1.76'))

    def test_rejects_non_xlsx_and_oversized_upload(self):
        with self.assertRaisesRegex(KpiImportError, 'Only .xlsx'):
            _parse_workbook(upload('tracker.xls', b'not xlsx'))
        too_large = upload('large.xlsx', b'x' * (10 * 1024 * 1024 + 1))
        with self.assertRaisesRegex(KpiImportError, '10 MB'):
            _parse_workbook(too_large)

    def test_tolerates_missing_ranges_formula_errors_and_preserves_zero(self):
        def introduce_quirks(workbook):
            december = workbook['Dec 2025']
            december['A3'] = None

            january = workbook['Jan 2026']
            january['A3'] = None
            january['D4'] = 0
            january['D5'] = '#VALUE!'
            january['D6'] = None

            february = workbook['Feb 2026']
            february['A3'] = '2/?? malformed'

            march = workbook['Mar 2026']
            for row in range(1, march.max_row + 1):
                if str(march.cell(row, 3).value or '').strip().lower() in {
                    'monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday', 'sunday',
                }:
                    march.cell(row, 4).value = None
                if str(march.cell(row, 2).value or '').strip().lower() == 'mtd':
                    march.cell(row, 4).value = 0

        content = mutate_tracker(tracker_bytes(2026), introduce_quirks)
        parsed = _parse_workbook(upload('quirks.xlsx', content))
        by_date = {row['business_date']: row for row in parsed.days}
        march = next(row for row in parsed.periods if row['sheet_name'] == 'Mar 2026')

        self.assertEqual(len(parsed.days), 364)
        self.assertEqual(by_date[date(2025, 12, 29)]['sales'], 0)
        self.assertIsNone(by_date[date(2025, 12, 30)]['sales'])
        self.assertIsNone(by_date[date(2025, 12, 31)]['sales'])
        self.assertIsNone(march['mtd_sales'])
        self.assertTrue(any('date range could not be read' in warning for warning in parsed.warnings))

    def test_rejects_unrecognized_required_header(self):
        def remove_header(workbook):
            workbook['Dec 2025']['B1'] = 'Unknown target column'

        content = mutate_tracker(tracker_bytes(2026), remove_header)
        with self.assertRaisesRegex(KpiImportError, 'missing required KPI headers'):
            _parse_workbook(upload('wrong-header.xlsx', content))


class KpiImportAndStateApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = get_user_model().objects.create_user(username='kpi-member', password='pw')
        self.organization = Organization.objects.create(name='KPI Store One')
        OrganizationMembership.objects.create(user=self.user, organization=self.organization)
        self.other_organization = Organization.objects.create(name='KPI Store Two')
        OrganizationMembership.objects.create(user=self.user, organization=self.other_organization)
        self.admin = get_user_model().objects.create_user(username='kpi-admin', password='pw', is_staff=True)
        self.import_url = f'/api/admin/organizations/{self.organization.id}/kpi-imports/'
        self.date = '2026-09-14'
        self.client.force_authenticate(self.user)
        self.client.credentials(HTTP_X_ORGANIZATION_ID=str(self.organization.id))

    def import_pair(self, *, target=12918):
        prior_metrics = {
            '2025-09-15': {
                'sales': 9522, 'traffic': 628, 'conversion': 0.172,
                'atv': 88.17, 'upt': 1.76,
            },
        }
        return self.client.post(
            self.import_url,
            {
                'current_year_file': upload('current.xlsx', tracker_bytes(2026, target=target)),
                'prior_year_file': upload('prior.xlsx', tracker_bytes(2025, prior_metrics=prior_metrics)),
            },
            format='multipart',
        )

    def patch_day(self, organization, payload):
        self.client.credentials(HTTP_X_ORGANIZATION_ID=str(organization.id))
        return self.client.patch(f'/api/schedule/kpi-days/{self.date}/', payload, format='json')

    def workbook_kpi(self, organization):
        self.client.credentials(HTTP_X_ORGANIZATION_ID=str(organization.id))
        response = self.client.get('/api/schedule/workbook/', {'week_start': '2026-09-14', 'day': 'Mon'})
        self.assertEqual(response.status_code, 200)
        return response.data['kpi']

    def test_member_cannot_import_but_platform_admin_can_and_history_has_hashes(self):
        self.assertEqual(self.import_pair().status_code, 403)

        self.client.force_authenticate(self.admin)
        response = self.import_pair()
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['current_fiscal_year'], 2026)
        self.assertEqual(response.data['prior_fiscal_year'], 2025)
        self.assertEqual(response.data['daily_records_imported'], 728)
        self.assertEqual(len(response.data['current_sha256']), 64)
        self.assertEqual(len(response.data['prior_sha256']), 64)
        self.assertEqual(KpiImportBatch.objects.count(), 1)
        self.assertEqual(KpiDailyRecord.objects.filter(organization=self.organization).count(), 728)
        self.assertEqual(KpiPeriodRecord.objects.filter(organization=self.organization).count(), 24)

        history = self.client.get(self.import_url)
        self.assertEqual(history.status_code, 200)
        self.assertEqual(history.data[0]['id'], response.data['id'])
        self.assertEqual(history.data[0]['imported_by']['username'], self.admin.username)

    def test_current_workbook_can_be_imported_alone_without_replacing_prior_year(self):
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.import_pair().status_code, 201)
        prior_record = KpiDailyRecord.objects.get(
            organization=self.organization, business_date=date(2025, 9, 15),
        )
        prior_batch_id = prior_record.import_batch_id
        prior_sales = prior_record.sales

        response = self.client.post(
            self.import_url,
            {'current_year_file': upload('current-only.xlsx', tracker_bytes(2026, target=22222))},
            format='multipart',
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['current_fiscal_year'], 2026)
        self.assertIsNone(response.data['prior_fiscal_year'])
        self.assertEqual(response.data['prior_filename'], '')
        self.assertEqual(response.data['daily_records_imported'], 364)
        self.assertEqual(KpiDailyRecord.objects.filter(organization=self.organization).count(), 728)
        preserved_prior = KpiDailyRecord.objects.get(
            organization=self.organization, business_date=date(2025, 9, 15),
        )
        self.assertEqual(preserved_prior.import_batch_id, prior_batch_id)
        self.assertEqual(preserved_prior.sales, prior_sales)
        self.assertEqual(
            KpiDailyRecord.objects.get(
                organization=self.organization, business_date=date(2026, 9, 14),
            ).day_target,
            Decimal('22222'),
        )

    def test_prior_workbook_can_be_imported_alone_without_replacing_current_year(self):
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.import_pair(target=13750).status_code, 201)
        current_record = KpiDailyRecord.objects.get(
            organization=self.organization, business_date=date(2026, 9, 14),
        )
        current_batch_id = current_record.import_batch_id

        response = self.client.post(
            self.import_url,
            {
                'prior_year_file': upload(
                    'prior-only.xlsx',
                    tracker_bytes(2025, prior_metrics={'2025-09-15': {'sales': 7777}}),
                ),
            },
            format='multipart',
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertIsNone(response.data['current_fiscal_year'])
        self.assertEqual(response.data['prior_fiscal_year'], 2025)
        self.assertEqual(response.data['current_filename'], '')
        self.assertEqual(response.data['daily_records_imported'], 364)
        self.assertEqual(KpiDailyRecord.objects.filter(organization=self.organization).count(), 728)
        preserved_current = KpiDailyRecord.objects.get(
            organization=self.organization, business_date=date(2026, 9, 14),
        )
        self.assertEqual(preserved_current.import_batch_id, current_batch_id)
        self.assertEqual(preserved_current.day_target, Decimal('13750'))
        self.assertEqual(
            KpiDailyRecord.objects.get(
                organization=self.organization, business_date=date(2025, 9, 15),
            ).sales,
            Decimal('7777'),
        )

    def test_import_requires_at_least_one_workbook(self):
        self.client.force_authenticate(self.admin)
        response = self.client.post(self.import_url, {}, format='multipart')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data['error'], 'Select at least one fiscal-year workbook to upload.')
        self.assertFalse(KpiImportBatch.objects.exists())

    def test_reversed_years_roll_back_without_replacing_active_rows(self):
        self.client.force_authenticate(self.admin)
        first = self.import_pair()
        self.assertEqual(first.status_code, 201, first.data)
        original_batch = KpiDailyRecord.objects.filter(
            organization=self.organization, business_date=date(2026, 9, 14),
        ).values_list('import_batch_id', flat=True).get()

        reversed_response = self.client.post(
            self.import_url,
            {
                'current_year_file': upload('current.xlsx', tracker_bytes(2025)),
                'prior_year_file': upload('prior.xlsx', tracker_bytes(2026)),
            },
            format='multipart',
        )

        self.assertEqual(reversed_response.status_code, 400)
        self.assertEqual(KpiImportBatch.objects.count(), 1)
        self.assertEqual(KpiDailyRecord.objects.filter(organization=self.organization).count(), 728)
        self.assertEqual(
            KpiDailyRecord.objects.get(
                organization=self.organization, business_date=date(2026, 9, 14),
            ).import_batch_id,
            original_batch,
        )

    def test_import_populates_workbook_comparison_and_month_metrics(self):
        self.client.force_authenticate(self.admin)
        response = self.import_pair()
        self.assertEqual(response.status_code, 201, response.data)
        self.client.force_authenticate(self.user)
        kpi = self.workbook_kpi(self.organization)

        self.assertEqual(kpi['comparisonDate'], '2025-09-15')
        self.assertEqual(kpi['goals']['daySalesTarget'], 12918)
        self.assertEqual(kpi['goals']['lastYearSales'], 9522)
        self.assertEqual(kpi['goals']['lastYearTraffic'], 628)
        self.assertEqual(kpi['goals']['conversionTarget'], 17.2)
        self.assertEqual(kpi['goals']['atv'], 88.17)
        self.assertEqual(kpi['goals']['upt'], 1.76)
        self.assertEqual(kpi['goals']['monthSalesPlan'], 652449)
        self.assertEqual(kpi['goals']['monthToDateSales'], 119829)

    def test_member_edits_are_shared_scoped_and_preserved_on_reimport(self):
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.import_pair().status_code, 201)
        self.client.force_authenticate(self.user)

        first = self.patch_day(self.organization, {
            'goal_updates': {'stretchTarget': 15000, 'trafficTrend': 10},
            'hourly_updates': {'10': {'actual': 1250, 'cel': 'AB'}},
        })
        self.assertEqual(first.status_code, 200, first.data)
        self.assertEqual(first.data['goals']['stretchTarget'], 15000)
        self.assertEqual(first.data['goals']['projectedTraffic'], 691)
        self.assertEqual(first.data['goals']['transactionGoal'], 119)
        self.assertEqual(first.data['hourly']['10']['actual'], 1250)
        self.assertEqual(first.data['hourly']['10']['cel'], 'AB')
        self.assertEqual(first.data['lastEditor']['username'], self.user.username)

        second = self.patch_day(self.other_organization, {'goal_updates': {'stretchTarget': 17000}})
        self.assertEqual(second.status_code, 200)
        self.assertIsNone(second.data['goals']['daySalesTarget'])
        self.assertEqual(second.data['goals']['stretchTarget'], 17000)
        self.assertEqual(self.workbook_kpi(self.organization)['goals']['stretchTarget'], 15000)

        self.client.force_authenticate(self.admin)
        self.assertEqual(self.import_pair(target=13000).status_code, 201)
        self.client.force_authenticate(self.user)
        refreshed = self.workbook_kpi(self.organization)
        self.assertEqual(refreshed['baseGoals']['daySalesTarget'], 13000)
        self.assertEqual(refreshed['goals']['stretchTarget'], 15000)
        self.assertEqual(refreshed['hourly']['10']['actual'], 1250)

        reset = self.patch_day(self.organization, {'goal_resets': ['stretchTarget']})
        self.assertEqual(reset.status_code, 200)
        self.assertIsNone(reset.data['goals']['stretchTarget'])
        self.assertEqual(reset.data['goals']['daySalesTarget'], 13000)

    def test_cross_organization_header_is_rejected_and_field_updates_merge(self):
        denied_org = Organization.objects.create(name='No Membership Store')
        self.client.credentials(HTTP_X_ORGANIZATION_ID=str(denied_org.id))
        denied = self.client.patch(
            f'/api/schedule/kpi-days/{self.date}/',
            {'goal_updates': {'stretchTarget': 15000}}, format='json',
        )
        self.assertEqual(denied.status_code, 403)

        one = self.patch_day(self.organization, {'goal_updates': {'stretchTarget': 15000}})
        two = self.patch_day(self.organization, {'goal_updates': {'trafficTrend': 5}})
        self.assertEqual(one.status_code, 200)
        self.assertEqual(two.status_code, 200)
        state = KpiDayState.objects.get(organization=self.organization, business_date=date(2026, 9, 14))
        self.assertEqual(state.goal_overrides, {'stretchTarget': 15000, 'trafficTrend': 5})
        self.assertEqual(state.revision, 2)
