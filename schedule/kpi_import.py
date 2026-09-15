import hashlib
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

import openpyxl
from django.db import transaction

from .models import KpiDailyRecord, KpiImportBatch, KpiPeriodRecord


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_MEMBERS = 5000
MAX_SHEETS = 24
MAX_ROWS_PER_SHEET = 200
MAX_COLUMNS_PER_SHEET = 100

MONTHS = {
    'jan': 1, 'january': 1, 'feb': 2, 'february': 2, 'mar': 3, 'march': 3,
    'apr': 4, 'april': 4, 'may': 5, 'jun': 6, 'june': 6, 'jul': 7,
    'july': 7, 'aug': 8, 'august': 8, 'sep': 9, 'sept': 9, 'september': 9,
    'oct': 10, 'october': 10, 'nov': 11, 'november': 11, 'dec': 12,
    'december': 12,
}
WEEKDAYS = {
    'monday': 0, 'tuesday': 1, 'wednesday': 2, 'thursday': 3,
    'friday': 4, 'saturday': 5, 'sunday': 6,
}
HEADER_ALIASES = {
    'day_target': {'incentivetargets', 'incentivetarget'},
    'day': {'day'},
    'sales': {'retailsalesnetvalue'},
    'conversion': {'conversionday'},
    'traffic': {'traffic'},
    'atv': {'retailsalesnetatvday'},
    'upt': {'retailsalesnetuptday'},
}


class KpiImportError(ValueError):
    pass


@dataclass
class UploadedWorkbook:
    name: str
    size: int
    sha256: str
    fiscal_year: int
    days: list
    periods: list
    warnings: list


def _normalize(value):
    return re.sub(r'[^a-z0-9]+', '', str(value or '').strip().lower())


def _decimal(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value or value.startswith('#'):
            return None
        value = value.replace(',', '').replace('$', '').replace('%', '')
    try:
        result = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return None
    return result if result.is_finite() else None


def _read_upload(upload):
    name = Path(upload.name or '').name
    if not name.lower().endswith('.xlsx'):
        raise KpiImportError('Only .xlsx files are supported.')
    if upload.size > MAX_UPLOAD_BYTES:
        raise KpiImportError('Each workbook must be 10 MB or smaller.')
    upload.seek(0)
    payload = upload.read()
    upload.seek(0)
    if len(payload) > MAX_UPLOAD_BYTES:
        raise KpiImportError('Each workbook must be 10 MB or smaller.')
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            members = archive.infolist()
            expanded = sum(member.file_size for member in members)
            if len(members) > MAX_ZIP_MEMBERS or expanded > MAX_UNCOMPRESSED_BYTES:
                raise KpiImportError('Workbook archive is too large to process safely.')
            if '[Content_Types].xml' not in archive.namelist():
                raise KpiImportError('File is not a valid XLSX workbook.')
            for member in members:
                compressed = max(member.compress_size, 1)
                if member.file_size > 1024 * 1024 and member.file_size / compressed > 1000:
                    raise KpiImportError('Workbook archive compression ratio is unsafe.')
    except zipfile.BadZipFile as exc:
        raise KpiImportError('File is not a valid XLSX workbook.') from exc
    return name, payload, hashlib.sha256(payload).hexdigest()


def _sheet_identity(title):
    normalized = title.strip().lower()
    if _normalize(normalized) == 'template':
        return None
    month = next((number for token, number in MONTHS.items() if re.search(rf'\b{token}\b', normalized)), None)
    year_match = re.search(r'\b(20\d{2})\b', normalized)
    if month is None or year_match is None:
        raise KpiImportError(f'Cannot identify month and year from sheet "{title}".')
    year = int(year_match.group(1))
    fiscal_period = 1 if month == 12 else month + 1
    return month, year, fiscal_period


def _find_headers(rows, sheet_name):
    aliases = {alias: key for key, values in HEADER_ALIASES.items() for alias in values}
    found = {}
    for row_index, row in enumerate(rows[:5]):
        for column_index, value in enumerate(row):
            key = aliases.get(_normalize(value))
            if key and key not in found:
                found[key] = column_index
    missing = [key for key in HEADER_ALIASES if key not in found]
    if missing:
        raise KpiImportError(f'Sheet "{sheet_name}" is missing required KPI headers: {", ".join(missing)}.')
    return found


def _valid_date(year, month, day):
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _range_start(value, sheet_month, sheet_year, expected=None):
    text = str(value or '').strip().lower()
    if not text:
        return None
    named = re.search(r'\b(' + '|'.join(sorted(MONTHS, key=len, reverse=True)) + r')\s*(\d{1,2})\b', text)
    candidates = []
    years = {sheet_year}
    if expected:
        years.update({expected.year - 1, expected.year, expected.year + 1})
    if named:
        month, day = MONTHS[named.group(1)], int(named.group(2))
        candidates = [candidate for year in years if (candidate := _valid_date(year, month, day))]
    else:
        numeric = re.search(r'(?<!\d)(\d{1,2})\s*/\s*(\d{1,2})(?!\d)', text)
        if not numeric:
            return None
        first, second = int(numeric.group(1)), int(numeric.group(2))
        pairs = []
        if first <= 12:
            pairs.append((first, second))
        if second <= 12 and (second, first) not in pairs:
            pairs.append((second, first))
        for year in years:
            candidates.extend(candidate for month, day in pairs if (candidate := _valid_date(year, month, day)))
    if not candidates:
        return None
    if expected:
        return min(candidates, key=lambda candidate: abs((candidate - expected).days))
    same_month = [candidate for candidate in candidates if candidate.month == sheet_month]
    return min(same_month or candidates)


def _parse_workbook(upload):
    filename, payload, sha256 = _read_upload(upload)
    try:
        workbook = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise KpiImportError('Workbook could not be opened as XLSX.') from exc
    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise KpiImportError(f'Workbook has more than {MAX_SHEETS} sheets.')
        sheets = []
        for worksheet in workbook.worksheets:
            identity = _sheet_identity(worksheet.title)
            if identity is None:
                continue
            if worksheet.max_row > MAX_ROWS_PER_SHEET or worksheet.max_column > MAX_COLUMNS_PER_SHEET:
                raise KpiImportError(f'Sheet "{worksheet.title}" exceeds supported dimensions.')
            values = [list(row) for row in worksheet.iter_rows(values_only=True)]
            headers = _find_headers(values, worksheet.title)
            sheets.append((identity[2], identity[0], identity[1], worksheet.title, values, headers))
        if len(sheets) != 12 or {sheet[0] for sheet in sheets} != set(range(1, 13)):
            raise KpiImportError('Workbook must contain one recognized sheet for each fiscal period from December through November.')
        sheets.sort(key=lambda sheet: sheet[0])
        fiscal_year = max(sheet[2] for sheet in sheets)
        blocks = []
        summaries = []
        warnings = []
        for period, month, sheet_year, sheet_name, rows, headers in sheets:
            all_anchors = []
            for row_index, row in enumerate(rows):
                raw = row[0] if row else None
                match = re.match(r'^wk(?:eek)?\s*([1-5])$', str(raw or '').strip(), re.IGNORECASE)
                if match:
                    all_anchors.append((row_index, int(match.group(1))))
            anchors = []
            for anchor in all_anchors:
                if anchor[1] == len(anchors) + 1:
                    anchors.append(anchor)
                if len(anchors) == 5:
                    break
            if not anchors:
                raise KpiImportError(f'Sheet "{sheet_name}" has no recognizable week blocks.')
            if len(all_anchors) > len(anchors):
                warnings.append(f'{sheet_name}: ignored copied week blocks below the main tracker table.')
            for anchor_index, (start_row, week_number) in enumerate(anchors):
                end_row = next(
                    (row_index for row_index, _ in all_anchors if row_index > start_row),
                    len(rows),
                )
                weekdays = {}
                range_label = None
                has_metrics = False
                for row_index in range(start_row, end_row):
                    row = rows[row_index]
                    column_a = row[0] if row else None
                    if range_label is None and isinstance(column_a, str) and re.search(r'\d\s*[/\-]', column_a):
                        range_label = column_a
                    day_value = row[headers['day']] if headers['day'] < len(row) else None
                    weekday = WEEKDAYS.get(str(day_value or '').strip().lower())
                    if weekday is None:
                        continue
                    if weekday in weekdays:
                        raise KpiImportError(f'Sheet "{sheet_name}" week {week_number} repeats a weekday.')
                    metrics = {
                        key: _decimal(row[column]) if column < len(row) else None
                        for key, column in headers.items() if key != 'day'
                    }
                    has_metrics = has_metrics or any(value is not None for value in metrics.values())
                    weekdays[weekday] = metrics
                active = has_metrics or range_label is not None
                if not active:
                    continue
                if set(weekdays) != set(range(7)):
                    raise KpiImportError(f'Sheet "{sheet_name}" week {week_number} must contain Monday through Sunday rows.')
                block = {
                    'period': period, 'month': month, 'sheet_year': sheet_year,
                    'sheet_name': sheet_name, 'week': week_number,
                    'range_label': range_label, 'weekdays': weekdays,
                }
                blocks.append(block)

            mtd_row = next((row for row in rows if len(row) > 1 and _normalize(row[1]) == 'mtd'), None)
            sales_plan = _decimal(mtd_row[2]) if mtd_row and len(mtd_row) > 2 else None
            mtd_sales = _decimal(mtd_row[3]) if mtd_row and len(mtd_row) > 3 else None
            summaries.append({
                'fiscal_year': fiscal_year, 'fiscal_period': period, 'sheet_name': sheet_name,
                'sales_plan': sales_plan, 'mtd_sales': mtd_sales,
            })
        if len(blocks) != 52:
            raise KpiImportError(f'Workbook contains {len(blocks)} active weeks; v1 requires exactly 52.')

        first_monday = None
        for block_index, block in enumerate(blocks):
            anchor = _range_start(block['range_label'], block['month'], block['sheet_year'])
            if anchor is not None and anchor.weekday() == 0:
                first_monday = anchor - timedelta(days=block_index * 7)
                break
        if first_monday is None:
            raise KpiImportError('Workbook needs at least one parseable Monday-start date range.')

        days = []
        for block_index, block in enumerate(blocks):
            monday = first_monday + timedelta(days=block_index * 7)
            label_start = _range_start(
                block['range_label'], block['month'], block['sheet_year'], expected=monday,
            )
            if block['range_label'] and label_start is None:
                warnings.append(f'{block["sheet_name"]} week {block["week"]}: date range could not be read; fiscal sequence used.')
            elif label_start and label_start != monday:
                warnings.append(f'{block["sheet_name"]} week {block["week"]}: date range differs from fiscal sequence; fiscal sequence used.')
            for weekday in range(7):
                metrics = block['weekdays'][weekday]
                days.append({
                    'fiscal_year': fiscal_year,
                    'fiscal_period': block['period'],
                    'fiscal_week': block['week'],
                    'weekday': weekday,
                    'business_date': monday + timedelta(days=weekday),
                    'day_target': metrics['day_target'],
                    'sales': metrics['sales'],
                    'traffic': int(metrics['traffic']) if metrics['traffic'] is not None else None,
                    'conversion': metrics['conversion'],
                    'atv': metrics['atv'],
                    'upt': metrics['upt'],
                })

        for summary in summaries:
            period_days = [day for day in days if day['fiscal_period'] == summary['fiscal_period']]
            actual_days = [day for day in period_days if day['sales'] is not None]
            summary['mtd_as_of'] = max((day['business_date'] for day in actual_days), default=None)
            if not actual_days:
                summary['mtd_sales'] = None

        missing = Counter()
        for day in days:
            for field in ('day_target', 'sales', 'traffic', 'conversion', 'atv', 'upt'):
                if day[field] is None:
                    missing[field] += 1
        for field, count in missing.items():
            warnings.append(f'FY{fiscal_year}: {count} daily {field.replace("_", " ")} value(s) unavailable.')
        for summary in summaries:
            if summary['sales_plan'] is None:
                warnings.append(f'{summary["sheet_name"]}: month sales plan unavailable.')

        return UploadedWorkbook(
            name=filename, size=len(payload), sha256=sha256, fiscal_year=fiscal_year,
            days=days, periods=summaries, warnings=warnings,
        )
    finally:
        workbook.close()


def _slot_set(parsed):
    return {
        (day['fiscal_period'], day['fiscal_week'], day['weekday'])
        for day in parsed.days
    }


def import_kpi_pair(organization, current_upload, prior_upload, user):
    current = _parse_workbook(current_upload)
    prior = _parse_workbook(prior_upload)
    if current.fiscal_year != prior.fiscal_year + 1:
        raise KpiImportError(
            f'Expected consecutive fiscal years; received FY{current.fiscal_year} and FY{prior.fiscal_year}.',
        )
    if _slot_set(current) != _slot_set(prior):
        raise KpiImportError('Current and prior workbooks do not contain matching fiscal week/day slots.')
    current_dates = {(d['fiscal_period'], d['fiscal_week'], d['weekday']): d['business_date'] for d in current.days}
    prior_dates = {(d['fiscal_period'], d['fiscal_week'], d['weekday']): d['business_date'] for d in prior.days}
    if any((current_dates[key] - prior_dates[key]).days != 364 for key in current_dates):
        raise KpiImportError('Workbook fiscal calendars are not aligned to a 364-day retail comparison.')

    warnings = (current.warnings + prior.warnings)[:200]
    with transaction.atomic():
        batch = KpiImportBatch.objects.create(
            organization=organization,
            current_fiscal_year=current.fiscal_year,
            prior_fiscal_year=prior.fiscal_year,
            current_filename=current.name,
            prior_filename=prior.name,
            current_size=current.size,
            prior_size=prior.size,
            current_sha256=current.sha256,
            prior_sha256=prior.sha256,
            daily_records_imported=len(current.days) + len(prior.days),
            period_records_imported=len(current.periods) + len(prior.periods),
            warnings=warnings,
            status=KpiImportBatch.STATUS_WARNINGS if warnings else KpiImportBatch.STATUS_COMPLETED,
            imported_by=user,
        )
        years = [current.fiscal_year, prior.fiscal_year]
        KpiDailyRecord.objects.filter(organization=organization, fiscal_year__in=years).delete()
        KpiPeriodRecord.objects.filter(organization=organization, fiscal_year__in=years).delete()
        KpiDailyRecord.objects.bulk_create([
            KpiDailyRecord(organization=organization, import_batch=batch, **day)
            for parsed in (prior, current) for day in parsed.days
        ])
        KpiPeriodRecord.objects.bulk_create([
            KpiPeriodRecord(organization=organization, import_batch=batch, **period)
            for parsed in (prior, current) for period in parsed.periods
        ])
    return batch


def serialize_import_batch(batch):
    imported_by = None
    if batch.imported_by:
        profile = getattr(batch.imported_by, 'profile', None)
        imported_by = {
            'id': batch.imported_by_id,
            'username': batch.imported_by.username,
            'full_name': profile.full_name if profile else batch.imported_by.get_full_name(),
        }
    return {
        'id': batch.id,
        'current_fiscal_year': batch.current_fiscal_year,
        'prior_fiscal_year': batch.prior_fiscal_year,
        'current_filename': batch.current_filename,
        'prior_filename': batch.prior_filename,
        'current_size': batch.current_size,
        'prior_size': batch.prior_size,
        'current_sha256': batch.current_sha256,
        'prior_sha256': batch.prior_sha256,
        'daily_records_imported': batch.daily_records_imported,
        'period_records_imported': batch.period_records_imported,
        'warnings': batch.warnings,
        'status': batch.status,
        'imported_by': imported_by,
        'imported_at': batch.imported_at.isoformat(),
    }
